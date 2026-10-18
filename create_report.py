from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from pathlib import Path
import json
root=Path(__file__).parent
stats={}
for m in ('philosophers-semaphore','philosophers-monitor','rw-semaphore','rw-monitor'):
    rows=[json.loads(l) for l in (root/'results'/f'{m}.jsonl').read_text().splitlines()]
    stats[m]=(len(rows),rows[-1]['ms'])
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
styles=doc.styles;normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(38,49,62);normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',23),('Heading 1',16),('Heading 2',12.5)]:
 st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=RGBColor(12,68,119)
def p(s='',style=None):return doc.add_paragraph(s,style)
def h(s,level=1):doc.add_heading(s,level)
def listp(s):p(s,'List Bullet')
def image(m,caption):
 doc.add_picture(str(root/'results'/(m+'.png')),width=Inches(6.7));q=p(caption);q.style='Caption'
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
 for r in rows:
  for c,v in zip(t.add_row().cells,r):c.text=str(v)
 return t
p('NMAM INSTITUTE OF TECHNOLOGY, NITTE').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('DEPARTMENT OF INFORMATION SCIENCE AND ENGINEERING').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('OPERATING SYSTEMS').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('PROGRAMMING ASSIGNMENT REPORT', 'Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Concurrency control using semaphores, monitors, and live Java–browser execution').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Student name: _____________________');p('USN: _____________________');p('Course: Operating Systems');p('Submitted to: _____________________');p('Date: _____________________')
p('Submission note: Complete the student information after reviewing the implementation and rerunning the programs locally.')
doc.add_page_break()
h('1. Problem Statement and Objectives')
h('1.1 Problem context',2)
p('Several threads may contend for the same resource. Without a coordination rule, one thread can change shared data while another reads it, or multiple threads can become stuck waiting for one another. This assignment demonstrates both problems using four independently selectable Java implementations. [1,2]')
h('1.2 Objectives',2)
for s in ['Represent four philosophers and four chopsticks as threads and shared synchronization resources.','Prevent deadlock during philosopher resource acquisition, using two different strategies.','Permit simultaneous readers but enforce exclusive writer access.','Use array-based resource mapping and an explicit FIFO request queue where appropriate.','Connect real Java execution to a browser visualization and preserve timestamped event logs.','Check safety invariants against actual recorded executions.']:listp(s)
h('2. Readers–Writers: Semaphore-Based Solution')
h('2.1 Design decision',2)
p('Five reader threads and three writer threads access one integer initialized to 10. Three fair Java semaphores serve separate roles: turnstile (entry fairness), roomEmpty (exclusive resource), and countLock (reader-count protection). The first active reader takes roomEmpty, and the last one releases it. A writer takes turnstile before roomEmpty, preventing newly arriving readers from continually overtaking it. [3]')
h('2.2 Algorithm and data structure',2)
p('Reader: pass the turnstile; update reader count under countLock; if first, acquire roomEmpty; read; decrement under countLock; if last, release roomEmpty. Writer: acquire turnstile, acquire roomEmpty, write, release roomEmpty, release turnstile. Storage includes Semaphore objects, an AtomicInteger for reader count, and an AtomicInteger for shared value.')
h('2.3 Expected properties',2)
p('Multiple readers can overlap. A writer cannot overlap any reader or another writer. The fair turnstile reduces writer starvation for this finite workload, although thread scheduling does not give hard execution-time guarantees.')
h('3. Readers–Writers: Monitor-Based Solution')
h('3.1 Design decision',2)
p('This variant does not reuse the semaphore routine. It uses synchronized methods, wait(), and notifyAll() around an ArrayDeque request queue. Waiting threads register identifiers such as R0 and W0. A reader can enter if no writer is active and there is no writer ahead of it in the queue. A writer can enter only at the head of the queue when the resource is completely free. [2,4]')
h('3.2 Queue-based admission',2)
p('The request queue adds a meaningful data-structure component. Reader requests before a writer may be admitted together; a later reader cannot bypass an already queued writer. Operations use Java intrinsic monitor locking to protect both the queue and active counts. These choices make ordering easy to explain during a viva.')
h('4. Dining Philosophers: Semaphore-Based Solution')
h('4.1 Mapping and deadlock avoidance',2)
p('Four philosopher threads share four fair binary semaphores stored in an array. For philosopher p, left = p and right = (p + 1) % 4. Even IDs acquire left then right. Odd IDs acquire right then left. This asymmetric ordering breaks a circular chain of waiting, so the classic all-left-hand deadlock is avoided. [1,3]')
h('4.2 Execution flow',2)
p('Think → become hungry → acquire first fork → acquire second fork → eat → release both forks → finish the round. Each philosopher performs three rounds. Non-neighboring philosophers may eat at the same time. Semaphore fairness on individual forks does not establish a formal starvation guarantee for the whole algorithm.')
h('5. Dining Philosophers: Monitor-Based Solution')
h('5.1 Atomic pair allocation',2)
p('A ForkMonitor keeps a boolean array inUse[4]. Its synchronized acquire method waits while either required fork is occupied. When both are free, the monitor marks them occupied in the same critical section. After eating, the thread marks both free and calls notifyAll(). The philosopher never holds one fork while waiting for the second. [2,4]')
h('5.2 Benefit and limitation',2)
p('This design demonstrates a contrasting policy: the semaphore version breaks circular wait by acquisition order, while the monitor version avoids partial allocations altogether. Both preserve safe shared-fork ownership. Neither variant claims strict fairness under unbounded workloads.')
h('6. Browser-Connected Execution and Visualization')
p('The application is a Java HTTP server using the JDK built-in HttpServer API. The browser serves HTML, CSS, and JavaScript from the same Java process. Clicking Start sends POST /api/run?mode=<selected-mode>. The Java backend starts actual threads and writes serialized, timestamped newline-delimited JSON events to the HTTP response. JavaScript reads the response stream with fetch(), updates thread/fork states, animates blue active resources, and appends the event timeline. An Export Log button produces CSV. No frontend animation engine pretends to run the algorithms independently.')
p('The interface uses black/navy backgrounds, electric-blue highlights, state cards, a circular philosopher table, and a shared-data display. Only one simulation executes at once to keep the log understandable.')
h('7. Actual Runs, Observations, and Evidence')
p('The Java 21 source compiled successfully using javac. Four POST requests to the local Java backend produced the recorded .jsonl files listed below. Measured times include intentional sleep delays used to make thread interaction observable; they do not compare raw algorithm speed.')
table(['Implementation','Events','Elapsed (ms)','Completed operations'],[(m,stats[m][0],stats[m][1],'12 eating rounds' if m.startswith('philosophers') else '5 reads, 3 writes') for m in stats])
p('Automated log verification passed for all four recorded runs: sequential event numbers, start/completion markers, non-overlapping neighboring eaters, and no simultaneous writer/reader or writer/writer activity.')
p('Image evidence note: Because the browser test environment restricted direct localhost navigation, the screenshots below were rendered from the original interface using replay of captured Java event traces. The trace files themselves are genuine Java-backend execution outputs. Rerun the app on your laptop and take live screenshots for final submission.')
for m,label in [('rw-semaphore','Readers–Writers with semaphore, event-log replay'),('rw-monitor','Readers–Writers with monitor, event-log replay'),('philosophers-semaphore','Dining Philosophers with semaphore, event-log replay'),('philosophers-monitor','Dining Philosophers with monitor, event-log replay')]:
 h('7. '+label.split(',')[0],2);image(m,'Screenshot: '+label+'.')
h('8. Synchronization Concepts Demonstrated')
for s in ['Critical section: code whose shared state must be protected.','Deadlock: a cycle of threads unable to proceed because each waits on resources held by another.','Starvation: a thread waits indefinitely while other requests keep getting served.','Semaphore: a permit-based object acquired and released explicitly.','Monitor: a synchronized object whose state and wait conditions are encapsulated together.','Thread: independently scheduled execution inside one Java process; threads share the arrays, queues, and counters.']:listp(s)
h('9. Semaphore versus Monitor')
table(['Criterion','Semaphore variants','Monitor variants'],[['Resource control','Explicit acquire/release','Synchronized methods and condition wait'],['Dining solution','Odd/even ordered binary semaphores','Atomic allocation of both forks'],['Readers–Writers solution','Fair turnstile + shared reader count','FIFO queue + admission predicates'],['Main data structures','Semaphore[4], counters','boolean[4], ArrayDeque requests'],['Fairness caveat','No absolute scheduling guarantee','Queued admission for RW; DP may starve']])
h('10. What Was Learned')
p('The demonstrations show why concurrency correctness depends on resource ownership and admission rules rather than on thread creation alone. A small design change—altering fork-acquisition order or enforcing atomic allocation—can remove a specific deadlock scenario. The reader/writer experiments also show why a counter must be protected, why an exclusive gate is needed for writes, and how a FIFO request queue makes fairness more explicit. The live Java-to-browser event stream makes synchronization behavior easier to inspect than a final console message.')
h('11. Assistance and Integrity Note')
p('This project was developed with AI-assisted drafting and coding support, followed by actual Java compilation, execution, and trace validation. The included code and narrative are custom to this design rather than copied from the example report. Before academic submission, the student should run the app personally, inspect the source, replace placeholder identity fields, capture local screenshots, and disclose assistance according to the institution’s requirements. No claim is made about Turnitin similarity or AI-detection outcomes.')
h('12. References')
for s in ['[1] Silberschatz, Galvin, and Gagne, Operating System Concepts, 10th ed., Wiley, 2018. https://codex.cs.yale.edu/avi/os-book/OS10/','[2] Oracle, Java Tutorials — Synchronization. https://docs.oracle.com/javase/tutorial/essential/concurrency/sync.html','[3] Oracle, Java SE — Semaphore API. https://docs.oracle.com/en/java/javase/21/docs/api/java.base/java/util/concurrent/Semaphore.html','[4] Oracle, Java Language Specification, Chapter 17: Threads and Locks. https://docs.oracle.com/javase/specs/jls/se21/html/jls-17.html','[5] Oracle, Java SE — HttpServer. https://docs.oracle.com/en/java/javase/21/docs/api/jdk.httpserver/com/sun/net/httpserver/HttpServer.html']:
 p(s)
h('13. Student Declaration')
p('I have reviewed the code and verification evidence, executed the work myself, and acknowledge all outside assistance and consulted sources as required by my institution.')
p('Name: ______________________    Signature: ______________________    Date: ______________')
out=root/'OS_Assignment_Report.docx';doc.save(out);print(out)
